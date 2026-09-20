# -*- coding: utf-8 -*-
from openpyxl import load_workbook

path = r"F:\CursorProject_Python\Bes TR.xlsm"
outp = r"F:\CursorProject_Python\_bes_tr_analysis\_sheet_dump.txt"
wb = load_workbook(path, read_only=True, data_only=True, keep_links=False)


def write_block(f, ws, title, min_row, max_row, min_col, max_col):
    f.write(f"\n\n======== {title} ========\n")
    for i, row in enumerate(
        ws.iter_rows(
            min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col, values_only=True
        ),
        min_row,
    ):
        cells = []
        for j, v in enumerate(row, min_col):
            if v is None:
                continue
            s = str(v).replace("\n", " / ")
            if not s.strip():
                continue
            if len(s) > 120:
                s = s[:120] + "..."
            cells.append(f"{j}:{s}")
        if cells:
            f.write(f"R{i} | " + " || ".join(cells) + "\n")


with open(outp, "w", encoding="utf-8") as f:
    write_block(f, wb["T1"], "T1 r1-40 c1-40", 1, 40, 1, 40)
    write_block(f, wb["T2"], "T2 r1-30 c1-25", 1, 30, 1, 25)
    write_block(f, wb["T5"], "T5 headers r1-8 c1-30", 1, 8, 1, 30)
    write_block(f, wb["T6"], "T6 r1-90 c1-6", 1, 90, 1, 6)
    write_block(f, wb["T7"], "T7 r1-40 c1-12", 1, 40, 1, 12)
    write_block(f, wb["T8"], "T8 ArrC r1-130 c1-8", 1, 130, 1, 8)
    write_block(f, wb["T9"], "T9 r1-12 c1-16", 1, 12, 1, 16)
    write_block(f, wb["W1"], "W1 r1-20 c1-12", 1, 20, 1, 12)
    write_block(f, wb["W4"], "W4 r1-20 c1-12", 1, 20, 1, 12)

wb.close()
print("wrote", outp)
