# -*- coding: utf-8 -*-
from openpyxl import load_workbook

path = r"F:\CursorProject_Python\Bes TR.xlsm"
outp = r"F:\CursorProject_Python\_bes_tr_analysis\_rowcounts.txt"
wb = load_workbook(path, read_only=True, data_only=True, keep_links=False)

def last_used(ws, col_idx, max_scan=None):
    last = 0
    nonempty = 0
    for i, row in enumerate(ws.iter_rows(min_col=col_idx, max_col=col_idx, values_only=True), 1):
        v = row[0]
        if v is not None and str(v).strip() != "":
            last = i
            nonempty += 1
        if max_scan and i >= max_scan:
            break
    return last, nonempty

with open(outp, "w", encoding="utf-8") as f:
    specs = {
        "T1": 4,
        "T2": 4,
        "T3": 2,
        "T4": 2,
        "T5": 3,
        "T6": 2,
        "T7": 3,
        "T8": 3,
        "T9": 3,
        "W1": 2,
        "W4": 3,
    }
    for name, col in specs.items():
        last, nonempty = last_used(wb[name], col)
        f.write(f"{name} col{col} last={last} nonempty={nonempty}\n")
        print(name, last, nonempty)

wb.close()
print("wrote", outp)
